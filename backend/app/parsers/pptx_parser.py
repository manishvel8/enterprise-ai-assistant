"""
pptx_parser.py — Parse PowerPoint presentations into ContentBlocks.

Library: python-pptx

Structure of a PPTX file:
  - Slides: each slide is a "page"
  - Each slide has shapes (text boxes, images, charts, tables)
  - Text boxes contain text frames → paragraphs → runs

Key fields we extract:
  - slide_number: which slide the content is from
  - slide_title: the title of the slide
  - content: body text from the slide
  - speaker_notes: notes below the slide

Why extract slide titles?
  Slide titles often summarize the content and are very useful as
  section_title metadata for the chunks — improves retrieval quality.
"""

import logging
from typing import List, Optional
from app.models.document import ContentBlock

logger = logging.getLogger(__name__)


def parse_pptx(file_bytes: bytes, document_id: str) -> List[ContentBlock]:
    """
    Parse a PPTX file and return ContentBlocks.

    One ContentBlock per slide (title + body text combined).
    Additional blocks for tables and speaker notes.
    """
    try:
        from pptx import Presentation
        from io import BytesIO
    except ImportError:
        logger.error("python-pptx not installed. Install: pip install python-pptx")
        return _fallback(document_id, "PPTX parser not available")

    blocks: List[ContentBlock] = []
    block_index = 0

    try:
        prs = Presentation(BytesIO(file_bytes))

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_title = _get_slide_title(slide)
            body_text_parts = []
            tables_found = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    shape_text = _extract_text_frame(shape)
                    if shape_text and shape_text != slide_title:
                        body_text_parts.append(shape_text)
                elif shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append("\t".join(cells))
                    tables_found.append("\n".join(rows))

            # Main slide block (title + body)
            slide_text_parts = []
            if slide_title:
                slide_text_parts.append(f"[Slide {slide_num}: {slide_title}]")
            slide_text_parts.extend(body_text_parts)
            slide_text = "\n".join(slide_text_parts).strip()

            if slide_text:
                blocks.append(ContentBlock(
                    block_id=f"{document_id}_block_{block_index}",
                    type="text",
                    text=slide_text,
                    slide_number=slide_num,
                    metadata={
                        "source": "pptx",
                        "slide_title": slide_title,
                    },
                ))
                block_index += 1

            # Table blocks
            for table_text in tables_found:
                if table_text.strip():
                    blocks.append(ContentBlock(
                        block_id=f"{document_id}_block_{block_index}",
                        type="table",
                        text=table_text,
                        slide_number=slide_num,
                        metadata={"source": "pptx_table", "slide_title": slide_title},
                    ))
                    block_index += 1

            # Speaker notes block
            notes_text = _get_speaker_notes(slide)
            if notes_text:
                blocks.append(ContentBlock(
                    block_id=f"{document_id}_block_{block_index}",
                    type="text",
                    text=f"[Speaker Notes - Slide {slide_num}] {notes_text}",
                    slide_number=slide_num,
                    metadata={"source": "pptx_notes", "slide_title": slide_title},
                ))
                block_index += 1

        logger.info(f"PPTX parsed: {len(prs.slides)} slides, {block_index} blocks")

    except Exception as e:
        logger.error(f"PPTX parsing failed: {e}")
        return _fallback(document_id, f"PPTX parsing error: {e}")

    return blocks


def _get_slide_title(slide) -> Optional[str]:
    """Extract the title from a slide (uses the title placeholder if available)."""
    try:
        if slide.shapes.title and slide.shapes.title.text:
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    return None


def _extract_text_frame(shape) -> str:
    """Extract all text from a shape's text frame."""
    parts = []
    for para in shape.text_frame.paragraphs:
        para_text = "".join(run.text for run in para.runs).strip()
        if para_text:
            parts.append(para_text)
    return "\n".join(parts).strip()


def _get_speaker_notes(slide) -> Optional[str]:
    """Extract speaker notes from a slide."""
    try:
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            text = notes_frame.text.strip() if notes_frame else ""
            return text if text else None
    except Exception:
        pass
    return None


def _fallback(document_id: str, message: str) -> List[ContentBlock]:
    return [ContentBlock(
        block_id=f"{document_id}_block_0",
        type="text",
        text=message,
        metadata={"source": "error"},
    )]
