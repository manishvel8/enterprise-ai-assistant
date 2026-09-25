"""
excel_parser.py — Parse Excel files (.xlsx) into ContentBlocks.

Library: openpyxl

Structure:
  - Workbook → Sheets → Rows → Cells
  - Each sheet becomes a ContentBlock
  - We serialize rows as tab-separated values (like CSV)

Why one block per sheet?
  - Different sheets often contain different topics (Q1 data, Q2 data, etc.)
  - The sheet name becomes the section_title for better retrieval context
"""

import logging
from typing import List
from app.models.document import ContentBlock

logger = logging.getLogger(__name__)


def parse_excel(file_bytes: bytes, document_id: str) -> List[ContentBlock]:
    """
    Parse an Excel .xlsx file and return ContentBlocks.

    One ContentBlock per sheet.
    Each cell row becomes a tab-separated line in the block text.
    Empty rows are skipped.
    """
    try:
        import openpyxl
        from io import BytesIO
    except ImportError:
        logger.error("openpyxl not installed. Install: pip install openpyxl")
        return _fallback(document_id, "Excel parser not available")

    blocks: List[ContentBlock] = []
    block_index = 0

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []

            # Get header row first (for context in chunk metadata)
            header_row = None
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    cells = [str(c) if c is not None else "" for c in row]
                    if header_row is None:
                        header_row = cells
                    rows.append("\t".join(cells))

            if rows:
                sheet_text = f"[Sheet: {sheet_name}]\n" + "\n".join(rows)
                blocks.append(ContentBlock(
                    block_id=f"{document_id}_block_{block_index}",
                    type="table",
                    text=sheet_text,
                    sheet_name=sheet_name,
                    metadata={
                        "source": "excel",
                        "sheet_name": sheet_name,
                        "row_count": len(rows),
                        "header": header_row,
                    },
                ))
                block_index += 1

        wb.close()
        logger.info(f"Excel parsed: {len(wb.sheetnames)} sheets, {block_index} blocks")

    except Exception as e:
        logger.error(f"Excel parsing failed: {e}")
        return _fallback(document_id, f"Excel parsing error: {e}")

    return blocks


def _fallback(document_id: str, message: str) -> List[ContentBlock]:
    return [ContentBlock(
        block_id=f"{document_id}_block_0",
        type="table",
        text=message,
        metadata={"source": "error"},
    )]
